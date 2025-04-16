import sqlite3
from docx import Document
from flask import Flask, render_template, request, redirect, send_file
from io import BytesIO
import re
import logging

# Налаштування логування для дебагу
logging.basicConfig(level=logging.DEBUG)
logger = logging.getLogger(__name__)

app = Flask(__name__)

# Підключення до бази даних
def connect_db():
    conn = sqlite3.connect('military_data.db')
    return conn

# Створення таблиць
def create_tables():
    conn = connect_db()
    cursor = conn.cursor()
    tables = [
        'unit_name',
        'position_rank_name',
        'subunit',
        'weapon_ammo_duty',
        'work_duration'
    ]
    for table in tables:
        cursor.execute(f"PRAGMA table_info({table})")
        columns = [col[1] for col in cursor.fetchall()]
        if not columns or 'value' not in columns:
            cursor.execute(f'DROP TABLE IF EXISTS {table}')
            cursor.execute(f'''
                CREATE TABLE {table} (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    value TEXT NOT NULL
                )
            ''')

    cursor.execute("PRAGMA table_info(commander)")
    columns = [col[1] for col in cursor.fetchall()]
    if not columns or 'rank' not in columns or 'name' not in columns or 'unit_id' not in columns:
        cursor.execute('DROP TABLE IF EXISTS commander')
        cursor.execute('''
            CREATE TABLE commander (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                rank TEXT NOT NULL,
                name TEXT NOT NULL,
                unit_id INTEGER NOT NULL,
                FOREIGN KEY (unit_id) REFERENCES unit_name(id)
            )
        ''')

    cursor.execute('SELECT COUNT(*) FROM unit_name')
    if cursor.fetchone()[0] == 0:
        cursor.execute('INSERT INTO unit_name (value) VALUES (?)', ('А0000',))
    cursor.execute('SELECT COUNT(*) FROM commander')
    if cursor.fetchone()[0] == 0:
        cursor.execute('INSERT INTO commander (rank, name, unit_id) VALUES (?, ?, ?)', ('Полковник', 'Петро Сидоров', 1))

    conn.commit()
    conn.close()

# Отримання даних з таблиці
def get_table_data(table_name):
    conn = connect_db()
    cursor = conn.cursor()
    if table_name == 'commander':
        cursor.execute('''
            SELECT commander.id, commander.rank, commander.name, unit_name.value 
            FROM commander 
            JOIN unit_name ON commander.unit_id = unit_name.id
        ''')
    else:
        cursor.execute(f'SELECT id, "value" FROM {table_name}')
    data = cursor.fetchall()
    conn.close()
    return data

# Отримання значення за ID
def get_value_by_id(table_name, id):
    if not id:
        logger.warning(f"ID is None for table {table_name}")
        return '' if table_name != 'commander' else ('', '', '')
    
    conn = connect_db()
    cursor = conn.cursor()
    if table_name == 'commander':
        cursor.execute('''
            SELECT commander.rank, commander.name, unit_name.value 
            FROM commander 
            JOIN unit_name ON commander.unit_id = unit_name.id 
            WHERE commander.id=?
        ''', (id,))
        result = cursor.fetchone()
        conn.close()
        logger.debug(f"Commander data for id {id}: {result}")
        return result if result else ('', '', '')
    else:
        cursor.execute(f'SELECT "value" FROM {table_name} WHERE id=?', (id,))
        result = cursor.fetchone()
        conn.close()
        logger.debug(f"Value for table {table_name}, id {id}: {result}")
        return result[0] if result else ''

# Додавання або оновлення даних
def upsert_data(table_name, id, value):
    conn = connect_db()
    cursor = conn.cursor()
    if table_name == 'commander':
        rank = value['rank']
        name = value['name']
        unit_id = value['unit_id']
        if id:
            cursor.execute('UPDATE commander SET rank=?, name=?, unit_id=? WHERE id=?', (rank, name, unit_id, id))
        else:
            cursor.execute('INSERT INTO commander (rank, name, unit_id) VALUES (?, ?, ?)', (rank, name, unit_id))
    else:
        if id:
            cursor.execute(f'UPDATE {table_name} SET "value"=? WHERE id=?', (value, id))
        else:
            cursor.execute(f'INSERT INTO {table_name} ("value") VALUES (?)', (value,))
    conn.commit()
    conn.close()

# Видалення запису
def delete_data(table_name, id):
    conn = connect_db()
    cursor = conn.cursor()
    cursor.execute(f'DELETE FROM {table_name} WHERE id=?', (id,))
    conn.commit()
    conn.close()

# Перетворення імені таблиці на зрозумілу назву
def table_name_to_display_name(table_name):
    return {
        'unit_name': 'Назва військової частини',
        'position_rank_name': 'Займана посада, звання, ПІ',
        'subunit': 'Підрозділ',
        'weapon_ammo_duty': 'Вид зброї, боєприпаси, особи наряду',
        'work_duration': 'Тривалість, година, куди, в чиє розпорядження',
        'commander': 'Звання, ім’я командира та назва частини'
    }[table_name]

# Функція для заміни плейсхолдерів у параграфі
def replace_placeholder_in_paragraph(paragraph, placeholder, replacement):
    full_text = ''.join(run.text for run in paragraph.runs).strip()
    logger.debug(f"Checking placeholder {placeholder} in text: {full_text}")
    if placeholder in full_text:
        new_text = full_text.replace(placeholder, replacement)
        runs = paragraph.runs
        if runs:
            first_run = runs[0]
            font = first_run.font
            bold = first_run.bold
            italic = first_run.italic
            underline = first_run.underline

            paragraph.clear()
            new_run = paragraph.add_run(new_text)
            new_run.bold = bold
            new_run.italic = italic
            new_run.underline = underline
            if font:
                new_run.font.name = font.name
                new_run.font.size = font.size
        logger.debug(f"Replaced {placeholder} with {replacement}")
        return True
    # Додаткова заміна для 00.00.0000, якщо плейсхолдер {date} не знайдено
    if placeholder == '{date}' and '00.00.0000' in full_text:
        new_text = full_text.replace('00.00.0000', replacement)
        runs = paragraph.runs
        if runs:
            first_run = runs[0]
            font = first_run.font
            bold = first_run.bold
            italic = first_run.italic
            underline = first_run.underline

            paragraph.clear()
            new_run = paragraph.add_run(new_text)
            new_run.bold = bold
            new_run.italic = italic
            new_run.underline = underline
            if font:
                new_run.font.name = font.name
                new_run.font.size = font.size
        logger.debug(f"Replaced 00.00.0000 with {replacement}")
        return True
    return False

# Функція для заміни номера пункту
def replace_number_in_paragraph(paragraph, old_number, new_number):
    full_text = ''.join(run.text for run in paragraph.runs).strip()
    logger.debug(f"Replacing number {old_number} with {new_number} in text: {full_text}")
    if old_number in full_text:
        new_text = full_text.replace(old_number, new_number)
        runs = paragraph.runs
        if runs:
            first_run = runs[0]
            font = first_run.font
            bold = first_run.bold
            italic = first_run.italic
            underline = first_run.underline

            paragraph.clear()
            new_run = paragraph.add_run(new_text)
            new_run.bold = bold
            new_run.italic = italic
            new_run.underline = underline
            if font:
                new_run.font.name = font.name
                new_run.font.size = font.size
        logger.debug(f"Number replaced: {old_number} -> {new_number}")
        return True
    return False

# Функція для створення документа
def create_order_document(date, order_number, unit_id, include_items, selected_values):
    doc = Document('template.docx')

    # Заміна загальних даних
    for para in doc.paragraphs:
        replace_placeholder_in_paragraph(para, '{date}', date)
        replace_placeholder_in_paragraph(para, '{order_number}', order_number)
        unit_name = get_value_by_id('unit_name', unit_id)
        replace_placeholder_in_paragraph(para, '{unit_id}', unit_name)

    # Заміна даних командира
    commander_rank, commander_name, commander_unit = get_value_by_id('commander', selected_values['commander'])
    commander_full = f"{commander_rank} {commander_name}" if commander_rank and commander_name else ''
    for para in doc.paragraphs:
        replace_placeholder_in_paragraph(para, '{commander}', commander_full)
        replace_placeholder_in_paragraph(para, '{rank}', commander_rank if commander_rank else '')
        replace_placeholder_in_paragraph(para, '{name}', commander_name if commander_name else '')
        replace_placeholder_in_paragraph(para, '{commander_unit}', commander_unit if commander_unit else '')

    # Мапінг плейсхолдерів
    placeholder_map = [
        ('{duty_officer}', 'position_rank_name', 'duty_officer', '1.1.'),
        ('{assistant_duty}', 'position_rank_name', 'assistant_duty', '1.2.'),
        ('{guard_chief}', 'position_rank_name', 'guard_chief', '1.3.'),
        ('{park_duty}', 'position_rank_name', 'park_duty', '1.4.'),
        ('{duty_subunit}', 'subunit', 'duty_subunit', '1.5.'),
        ('{subunit_duty}', 'subunit', 'subunit_duty', '1.6.'),
        ('{work_duty}', 'work_duration', 'work_duty', '1.7.'),
        ('{weapon_ammo}', 'weapon_ammo_duty', 'weapon_ammo', '1.8.')
    ]

    # Збираємо активні пункти
    active_items = []
    paragraphs_to_remove = []
    for para in doc.paragraphs:
        full_text = ''.join(run.text for run in para.runs).strip()
        match = re.match(r'1\.(\d+)\.\s*', full_text)
        if match:
            old_number = match.group(0).strip()
            for placeholder, table, key, expected_number in placeholder_map:
                if placeholder in full_text and old_number == expected_number:
                    if include_items.get(key, False) and selected_values.get(key):
                        value = get_value_by_id(table, selected_values[key])
                        if value and value.strip():
                            active_items.append((para, old_number, placeholder, value))
                        else:
                            paragraphs_to_remove.append(para)
                    else:
                        paragraphs_to_remove.append(para)
                    break

    # Видаляємо неактивні пункти
    for para in paragraphs_to_remove:
        logger.debug(f"Removing paragraph: {para.text}")
        para._element.getparent().remove(para._element)

    # Замінюємо плейсхолдери та перенумеровуємо
    current_number = 1
    for para, old_number, placeholder, value in active_items:
        replace_placeholder_in_paragraph(para, placeholder, value)
        full_text = ''.join(run.text for run in para.runs).strip()
        if not full_text or full_text == old_number:
            logger.debug(f"Removing empty paragraph after replacement: {full_text}")
            para._element.getparent().remove(para._element)
            continue
        new_number = f"1.{current_number}."
        replace_number_in_paragraph(para, old_number, new_number)
        current_number += 1

    # Збереження документа
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# Головна сторінка
@app.route('/')
def index():
    tables = [
        ('unit_name', 'Назва військової частини'),
        ('position_rank_name', 'Займана посада, військове звання, прізвище та ініціали'),
        ('subunit', 'Підрозділ'),
        ('weapon_ammo_duty', 'Вид зброї, кількість боєприпасів і посадові особи'),
        ('work_duration', 'Тривалість роботи, година, куди та в чиє розпорядження'),
        ('commander', 'Звання, ім’я командира та назва частини')
    ]
    return render_template('index.html', tables=tables)

# Сторінка таблиці
@app.route('/table/<table_name>', methods=['GET', 'POST'])
def table_view(table_name):
    if request.method == 'POST':
        if table_name == 'commander':
            rank = request.form['rank']
            name = request.form['name']
            unit_id = request.form['unit_id']
            id = request.form.get('id')
            upsert_data(table_name, id, {'rank': rank, 'name': name, 'unit_id': unit_id})
        else:
            value = request.form['value']
            id = request.form.get('id')
            upsert_data(table_name, id, value)
        return redirect(f'/table/{table_name}')
    
    data = get_table_data(table_name)
    column_name = table_name_to_display_name(table_name)
    unit_names = get_table_data('unit_name') if table_name == 'commander' else []
    return render_template('table.html', table_name=table_name, data=data, column_name=column_name, unit_names=unit_names)

# Видалення запису
@app.route('/delete/<table_name>/<int:id>')
def delete(table_name, id):
    delete_data(table_name, id)
    return redirect(f'/table/{table_name}')

# Сторінка генерації документа
@app.route('/generate', methods=['GET', 'POST'])
def generate():
    if request.method == 'POST':
        date = request.form['date']
        order_number = request.form['order_number']
        unit_id = request.form['unit_id']
        include_items = {
            'duty_officer': 'duty_officer' in request.form,
            'assistant_duty': 'assistant_duty' in request.form,
            'guard_chief': 'guard_chief' in request.form,
            'park_duty': 'park_duty' in request.form,
            'duty_subunit': 'duty_subunit' in request.form,
            'subunit_duty': 'subunit_duty' in request.form,
            'work_duty': 'work_duty' in request.form,
            'weapon_ammo': 'weapon_ammo' in request.form
        }
        selected_values = {
            'duty_officer': request.form.get('duty_officer_id'),
            'assistant_duty': request.form.get('assistant_duty_id'),
            'guard_chief': request.form.get('guard_chief_id'),
            'park_duty': request.form.get('park_duty_id'),
            'duty_subunit': request.form.get('duty_subunit_id'),
            'subunit_duty': request.form.get('subunit_duty_id'),
            'work_duty': request.form.get('work_duty_id'),
            'weapon_ammo': request.form.get('weapon_ammo_id'),
            'commander': request.form.get('commander_id', '1')
        }
        logger.debug(f"Include items: {include_items}")
        logger.debug(f"Selected values: {selected_values}")

        buffer = create_order_document(date, order_number, unit_id, include_items, selected_values)
        return send_file(
            buffer,
            as_attachment=True,
            download_name=f"order_{date}.docx",
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

    unit_data = get_table_data('unit_name')
    position_rank_data = get_table_data('position_rank_name')
    subunit_data = get_table_data('subunit')
    work_duration_data = get_table_data('work_duration')
    weapon_ammo_data = get_table_data('weapon_ammo_duty')
    commander_data = get_table_data('commander')
    
    return render_template('generate.html', 
                         unit_data=unit_data,
                         position_rank_data=position_rank_data,
                         subunit_data=subunit_data,
                         work_duration_data=work_duration_data,
                         weapon_ammo_data=weapon_ammo_data,
                         commander_data=commander_data)

# Ініціалізація
if __name__ == '__main__':
    create_tables()
    app.run(debug=True, host='0.0.0.0', port=5000)